"""
Генерация Excel-отчётов по алертам дефектуры.

Отчёт содержит:
  - Топ-N продуктов под риском OOS с диагностическими признаками
  - Цветовое выделение строк по уровню риска (критический/высокий/средний/низкий)
  - Русскоязычные заголовки колонок
  - Заморозку строки заголовка
  - Автоширину колонок

Функции:
  save_excel_report()      — основной отчёт (один день)
  generate_summary_report() — сводный отчёт по нескольким датам
"""

import os
from pathlib import Path
from typing import Optional

import pandas as pd
from loguru import logger

try:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    logger.warning("openpyxl не установлен — форматирование Excel недоступно. "
                   "Установите: pip install openpyxl")


# ---------------------------------------------------------------------------
# Конфигурация колонок и цветов
# ---------------------------------------------------------------------------

# Hex-цвета без # (openpyxl)
COLORS = {
    "critical":    "FF9999",  # красный  (risk >= 0.80)
    "high":        "FFCC99",  # оранжевый (risk >= 0.60)
    "medium":      "FFFFCC",  # жёлтый  (risk >= 0.40)
    "low":         "CCFFCC",  # зелёный  (risk < 0.40)
    "header_bg":   "2F5496",  # тёмно-синий заголовок
    "header_font": "FFFFFF",  # белый шрифт заголовка
}

# Порядок колонок в отчёте
REPORT_COLUMNS = [
    "risk_rank",
    "code_kag",
    "risk_score",
    "is_alert",
    "total_stock",
    "n_zero",
    "days_of_stock",
    "runway_critical",
    "runway_warning",
    "n_red_flags",
    "critical_combination",
    "avg_sales_pct_7d",
    "avg_supply_pct",
    "supply_deficit",
    "persistent_outflow",
    "sales_velocity",
    "sales_accelerating",
    "oos_count_30d",
    "chronic_oos",
    "stock_zscore",
    "stock_anomaly_low",
    "consecutive_falls",
    "zeros_spreading",
    "prediction_date",
    "model_name",
]

# Русскоязычные заголовки
COLUMN_NAMES_RU = {
    "risk_rank":             "Ранг риска",
    "code_kag":              "КАГ-код",
    "risk_score":            "Оценка риска",
    "is_alert":              "Алерт",
    "total_stock":           "Общий остаток",
    "n_zero":                "Дистр. с нулём",
    "days_of_stock":         "Дней запаса",
    "runway_critical":       "Критич. остаток (<7д)",
    "runway_warning":        "Предупреждение (<14д)",
    "n_red_flags":           "Красных флагов",
    "critical_combination":  "Крит. комбинация",
    "avg_sales_pct_7d":      "Темп продаж 7д",
    "avg_supply_pct":        "Темп поставок",
    "supply_deficit":        "Дефицит поставок",
    "persistent_outflow":    "Устойчивый отток",
    "sales_velocity":        "Ускорение продаж",
    "sales_accelerating":    "Продажи растут",
    "oos_count_30d":         "OOS за 30д",
    "chronic_oos":           "Хроническая дефект.",
    "stock_zscore":          "Z-оценка остатка",
    "stock_anomaly_low":     "Аномально низкий",
    "consecutive_falls":     "Дней падения подряд",
    "zeros_spreading":       "Нули распространяются",
    "prediction_date":       "Дата прогноза",
    "model_name":            "Модель",
}


def _risk_level(score: float) -> str:
    if score >= 0.80:
        return "critical"
    if score >= 0.60:
        return "high"
    if score >= 0.40:
        return "medium"
    return "low"


# ---------------------------------------------------------------------------
# Основной отчёт (один день)
# ---------------------------------------------------------------------------

def save_excel_report(
    predictions: pd.DataFrame,
    date_str: str,
    reports_dir: str = "reports",
    filename: Optional[str] = None,
) -> str:
    """
    Сохранить Excel-отчёт с алертами дефектуры.

    Args:
        predictions: Результат функции predict() или batch_predict()
        date_str:    Дата отчёта в формате YYYY-MM-DD (используется в имени файла)
        reports_dir: Директория для сохранения (создаётся автоматически)
        filename:    Кастомное имя файла; по умолчанию oos_alerts_{date_str}.xlsx

    Returns:
        Полный путь к сохранённому файлу
    """
    Path(reports_dir).mkdir(parents=True, exist_ok=True)

    fname = filename or f"oos_alerts_{date_str}.xlsx"
    report_path = os.path.join(reports_dir, fname)

    # Выбираем и переименовываем колонки
    cols = [c for c in REPORT_COLUMNS if c in predictions.columns]
    df_report = predictions[cols].copy()
    df_report = df_report.rename(columns=COLUMN_NAMES_RU)

    # Форматирование числовых колонок
    for col_ru in [COLUMN_NAMES_RU.get("risk_score", "Оценка риска")]:
        if col_ru in df_report.columns:
            df_report[col_ru] = df_report[col_ru].round(4)

    if HAS_OPENPYXL:
        risk_scores = (
            predictions["risk_score"].values
            if "risk_score" in predictions.columns
            else None
        )
        _write_formatted_excel(df_report, risk_scores, report_path)
    else:
        df_report.to_excel(report_path, index=False)

    n_rows = len(df_report)
    n_alerts = int(predictions["is_alert"].sum()) if "is_alert" in predictions.columns else "?"
    logger.info(f"Отчёт сохранён: {report_path} ({n_rows} строк, алертов: {n_alerts})")
    return report_path


def _write_formatted_excel(
    df_report: pd.DataFrame,
    risk_scores,
    path: str,
) -> None:
    """Записать xlsx с форматированием через openpyxl."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df_report.to_excel(writer, index=False, sheet_name="Алерты")
        ws = writer.sheets["Алерты"]

        # Стиль заголовка
        header_fill = PatternFill("solid", fgColor=COLORS["header_bg"])
        header_font = Font(bold=True, color=COLORS["header_font"], size=10)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        ws.row_dimensions[1].height = 36

        # Автоширина колонок
        for col_idx, col_name in enumerate(df_report.columns, 1):
            header_len = len(str(col_name))
            max_data_len = (
                df_report[col_name].astype(str).str.len().max()
                if len(df_report) > 0 else 0
            )
            width = min(max(header_len, max_data_len) + 2, 40)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Цветовая заливка строк по уровню риска
        if risk_scores is not None:
            data_align = Alignment(horizontal="center", vertical="center")
            for row_idx, score in enumerate(risk_scores, start=2):
                level = _risk_level(float(score))
                fill = PatternFill("solid", fgColor=COLORS[level])
                for col_idx in range(1, len(df_report.columns) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.fill = fill
                    cell.alignment = data_align

        # Заморозка заголовка
        ws.freeze_panes = "A2"

        # Автофильтр
        ws.auto_filter.ref = ws.dimensions


# ---------------------------------------------------------------------------
# Сводный отчёт (несколько дат)
# ---------------------------------------------------------------------------

def generate_summary_report(
    all_predictions: pd.DataFrame,
    reports_dir: str = "reports",
) -> str:
    """
    Сгенерировать сводный Excel-отчёт по нескольким датам.

    Листы в файле:
      1. "Все прогнозы"        — все строки predictions
      2. "Статистика по датам" — агрегаты за каждую дату
      3. "Топ хронических"     — продукты с наибольшим средним риском

    Args:
        all_predictions: Объединённый DataFrame из batch_predict()
        reports_dir:     Директория для сохранения

    Returns:
        Путь к файлу summary_report.xlsx
    """
    Path(reports_dir).mkdir(parents=True, exist_ok=True)
    path = os.path.join(reports_dir, "summary_report.xlsx")

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # Лист 1: все прогнозы
        all_predictions.to_excel(writer, index=False, sheet_name="Все прогнозы")

        # Лист 2: статистика по датам
        if "prediction_date" in all_predictions.columns:
            daily = all_predictions.groupby("prediction_date").agg(
                продуктов=("code_kag", "nunique"),
                алертов=("is_alert", "sum"),
                средний_риск=("risk_score", "mean"),
                макс_риск=("risk_score", "max"),
            ).reset_index()
            daily.columns = ["Дата", "Продуктов", "Алертов", "Средний риск", "Макс риск"]
            daily.to_excel(writer, index=False, sheet_name="Статистика по датам")

        # Лист 3: хронические риски (топ-50 по среднему risk_score)
        if "risk_score" in all_predictions.columns and "code_kag" in all_predictions.columns:
            chronic = (
                all_predictions.groupby("code_kag")["risk_score"]
                .mean()
                .sort_values(ascending=False)
                .head(50)
                .reset_index()
            )
            chronic.columns = ["КАГ-код", "Средний риск"]
            chronic["Средний риск"] = chronic["Средний риск"].round(4)
            chronic.to_excel(writer, index=False, sheet_name="Топ хронических")

    logger.info(f"Сводный отчёт сохранён: {path}")
    return path
